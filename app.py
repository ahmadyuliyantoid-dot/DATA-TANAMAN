from functools import wraps
from flask import Flask, request, redirect, url_for, render_template_string, send_file, session
import sqlite3
import qrcode
import io
import os

app = Flask(__name__)
app.secret_key = os.environ.get("TANAMAN_SECRET_KEY", "a2af3cb252623fa8a713df56f5aa4b4c8c299bd97e6403d2b79e263fe37c04a0")
DB = r"D:\DATA_TANAMAN\tanaman.db"

def db():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c

with db() as c:
    c.execute("""
    CREATE TABLE IF NOT EXISTS tanaman(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nomor TEXT UNIQUE NOT NULL,
        nama TEXT NOT NULL,
        jenis TEXT,
        tanggal TEXT,
        lokasi TEXT,
        kondisi TEXT,
        tinggi TEXT,
        catatan TEXT
    )
    """)

STYLE = """
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>
body{font-family:Arial;background:#eef5ef;margin:0;color:#24372b}
header{background:#176b43;color:white;padding:22px}
main{max-width:1000px;margin:25px auto;padding:0 15px}
.card{background:white;padding:22px;border-radius:12px;margin-bottom:18px}
input,select,textarea{box-sizing:border-box;width:100%;padding:12px;margin:7px 0 14px;border:1px solid #ccd8ce;border-radius:6px}
button,.btn{display:inline-block;background:#176b43;color:white;padding:11px 15px;border:0;border-radius:6px;text-decoration:none;cursor:pointer}
.secondary{background:#e4eee7;color:#24372b}
.grid{display:grid;grid-template-columns:1fr 1fr;gap:12px}
.stats{display:flex;gap:12px;flex-wrap:wrap}
.stat{background:#e0f3e6;padding:18px;border-radius:8px;flex:1}
table{width:100%;border-collapse:collapse;min-width:570px}
th,td{text-align:left;padding:12px;border-bottom:1px solid #ddd}
th{background:#edf5ef}.scroll{overflow-x:auto}.qr{width:220px}
@media(max-width:600px){.grid{grid-template-columns:1fr}}
@media print{header,.no-print{display:none}}
</style>
"""

def page(title, body):
    return render_template_string("""
<!doctype html>
<html lang="id">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{{ title }} | DATA TANAMAN</title>
<link rel="stylesheet" href="/static/css/app.css?v=layout20261020">
</head>
<body>
<div class="layout-shell">
    <aside class="sidebar" id="sidebar">
        <a class="brand" href="/">
            <span class="brand-icon">🌿</span>
            <span>
                <strong>DATA TANAMAN</strong>
                <small>Plant Management System</small>
            </span>
        </a>

        <div class="menu-label">MENU UTAMA</div>
        <nav class="side-menu">
            <a href="/" class="menu-link">
                <span>▦</span> Dashboard & Data
            </a>
            <a href="/tambah" class="menu-link">
                <span>＋</span> Tambah Tanaman
            </a>
            <a href="/dashboard-lanjutan" class="menu-link">
                <span>▥</span> Analitik Tanaman
            </a>
        </nav>

        <div class="menu-label">MONITORING TANAMAN</div>
<nav class="side-menu">
    <a href="http://{{ request.host.split(":")[0] }}:5001/" class="menu-link"><span>◉</span> Monitoring Tanaman</a>
    <a href="http://{{ request.host.split(":")[0] }}:5003/" class="menu-link"><span>☑</span> Pemeriksaan Tanaman</a>
    <a href="http://{{ request.host.split(":")[0] }}:5006/" class="menu-link"><span>✿</span> Perawatan Tanaman</a>
    <a href="http://{{ request.host.split(":")[0] }}:5004/" class="menu-link"><span>▤</span> Riwayat Tanaman</a>
    <a href="http://{{ request.host.split(":")[0] }}:5009/" class="menu-link"><span>▥</span> Laporan Monitoring</a>
</nav>
<div class="menu-label">PENGELOLAAN</div>
        <nav class="side-menu">
            <a href="/#daftar-tanaman" class="menu-link">
                <span>▤</span> Daftar Tanaman
            </a>
            <a href="/login" class="menu-link">
                <span>♙</span> Login / Admin
            </a>
        </nav>

        <div class="sidebar-bottom">
            <div class="sidebar-note">
                <span>🌱</span>
                <strong>Inventaris Tanaman</strong>
                <small>Kelola data lebih mudah dan teratur.</small>
            </div>
        </div>
    </aside>

    <div class="main-area">
        <header class="topbar">
            <button class="menu-toggle" onclick="document.getElementById('sidebar').classList.toggle('sidebar-open')" aria-label="Buka menu">☰</button>
            <div class="topbar-title">
                <strong>{{ title }}</strong>
                <small>Sistem Manajemen Data Tanaman</small>
            </div>
            <a href="/tambah" class="topbar-action">＋ Tambah Data</a>
        </header>

        <main class="page-content">
            <div class="page-heading">
                <div>
                    <div class="eyebrow">PLANT INVENTORY SYSTEM</div>
                    <h1>{{ title }}</h1>
                    <p>Kelola inventaris dan informasi tanaman Anda.</p>
                </div>
            </div>
            <section class="page-body">
                {{ body|safe }}
            </section>
        </main>

        <footer class="app-footer">
            <span>DATA TANAMAN</span>
            <span>Sistem Inventaris Tanaman</span>
        </footer>
    </div>
</div>
</body>
</html>
""", title=title, body=body)

def admin_required(f):
    @wraps(f)
    def pembungkus(*args, **kwargs):
        if not session.get("admin"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return pembungkus


@app.route("/login", methods=["GET", "POST"])
def login():
    pesan = ""
    if request.method == "POST":
        nama = request.form.get("username", "")
        sandi = request.form.get("password", "")
        if nama == "admin" and sandi == "12345":
            session.clear()
            session["admin"] = True
            return redirect("/")
        pesan = "Username atau password salah."

    return page("Login Admin", render_template_string("""
    <div class="card" style="max-width:420px;margin:40px auto">
      <h2>Login Admin</h2>
      <p>{{pesan}}</p>
      <form method="post">
        <label>Username</label>
        <input name="username" required autocomplete="username">
        <label>Kata Sandi</label>
        <input name="password" type="password" required
               autocomplete="current-password">
        <button type="submit">Masuk</button>
        <a class="btn secondary" href="/">Kembali</a>
      </form>
    </div>
    """, pesan=pesan))


@app.route("/logout")
def logout():
    session.clear()
    return redirect("/")


@app.route("/")
def index():
    cari=request.args.get("cari","").strip()
    with db() as c:
        rows=c.execute("SELECT * FROM tanaman WHERE nomor LIKE ? OR nama LIKE ? OR lokasi LIKE ? ORDER BY id DESC",
            (f"%{cari}%",f"%{cari}%",f"%{cari}%")).fetchall()
        total=c.execute("SELECT COUNT(*) FROM tanaman").fetchone()[0]
        sehat=c.execute("SELECT COUNT(*) FROM tanaman WHERE kondisi='Sehat'").fetchone()[0]
    body=render_template_string("""
    <div class="stats"><div class="stat">Total Tanaman<h2>{{total}}</h2></div><div class="stat">Pohon Sehat<h2>{{sehat}}</h2></div></div>
    <div class="card">
    {% if session.get('admin') %}
      <a class="btn" href="/tambah">+ Tambah Pohon</a>
      <a class="btn secondary" href="/logout">Keluar Admin</a>
    {% else %}
      <a class="btn secondary" href="/login">Login Admin</a>
    {% endif %}
    </div>
    <div class="card no-print">
    <h2>Ekspor dan Cetak Laporan</h2>
    <p>Unduh seluruh data tanaman atau cetak daftar melalui browser.</p>
    <button type="button" onclick="window.print()">Print / Cetak</button>
    <a class="btn" href="/export/excel">Unduh Excel (.xlsx)</a>
    <a class="btn" href="/export/pdf">Unduh PDF (.pdf)</a>
    </div>
    <div class="card"><h2>Daftar Tanaman</h2>
    <form><input name="cari" placeholder="Cari nomor, nama, lokasi" value="{{cari}}"><button>Cari</button> <a class="btn secondary" href="/">Reset</a></form>
    <div class="scroll"><table><tr><th>Nomor</th><th>Nama</th><th>Lokasi</th><th>Kondisi</th><th>Detail</th></tr>
    {% for t in rows %}<tr><td>{{t['nomor']}}</td><td>{{t['nama']}}</td><td>{{t['lokasi'] or '-'}}</td><td>{{t['kondisi'] or '-'}}</td><td><a href="/tanaman/{{t['id']}}">Buka / QR</a></td></tr>
    {% else %}<tr><td colspan="5">Belum ada data tanaman. Klik Tambah Pohon.</td></tr>{% endfor %}
    </table></div></div>""",rows=rows,total=total,sehat=sehat,cari=cari)
    return page("Data Tanaman",body)


FORM = r'''
<div class="card" style="max-width:900px;margin:20px auto;padding:24px">
  <h2>{{ judul }}</h2>
  <p>Lengkapi informasi tanaman di bawah ini.</p>
  <form method="post">
    <div style="display:grid;grid-template-columns:repeat(auto-fit,minmax(230px,1fr));gap:18px">
      <div>
        <label>Nomor Tanaman</label>
        <input name="nomor" value="{{ nomor }}" required
          style="width:100%;box-sizing:border-box;padding:11px;margin-top:6px">
      </div>
      <div>
        <label>Nama Tanaman</label>
        <input name="nama" value="{{ t['nama'] if t else '' }}" required
          style="width:100%;box-sizing:border-box;padding:11px;margin-top:6px">
      </div>
      <div>
        <label>Jenis Tanaman</label>
        <input name="jenis" value="{{ t['jenis'] if t else '' }}"
          style="width:100%;box-sizing:border-box;padding:11px;margin-top:6px">
      </div>
      <div>
        <label>Tanggal Tanam</label>
        <input type="date" name="tanggal" value="{{ t['tanggal'] if t else '' }}"
          style="width:100%;box-sizing:border-box;padding:11px;margin-top:6px">
      </div>
      <div>
        <label>Lokasi</label>
        <input name="lokasi" value="{{ t['lokasi'] if t else '' }}"
          style="width:100%;box-sizing:border-box;padding:11px;margin-top:6px">
      </div>
      <div>
        <label>Kondisi Tanaman</label>
        <select name="kondisi"
          style="width:100%;box-sizing:border-box;padding:11px;margin-top:6px">
          {% for opsi in ['Sehat','Perlu Perawatan','Rusak','Mati'] %}
          <option value="{{ opsi }}" {% if t and t['kondisi']==opsi %}selected{% elif not t and opsi=='Sehat' %}selected{% endif %}>{{ opsi }}</option>
          {% endfor %}
        </select>
      </div>
      <div>
        <label>Catatan</label>
        <textarea name="catatan" rows="3"
          style="width:100%;box-sizing:border-box;padding:11px;margin-top:6px">{{ t['catatan'] if t else '' }}</textarea>
      </div>
    </div>
    <div style="display:flex;gap:10px;flex-wrap:wrap;margin-top:24px">
      <button type="submit" class="btn">Simpan Data</button>
      <a href="/" class="btn secondary">Kembali</a>
    </div>
  </form>
</div>
'''

@app.route("/tambah",methods=["GET","POST"])
@admin_required
def tambah():
    if request.method=="POST":
        f=request.form
        try:
            with db() as c:
                c.execute("INSERT INTO tanaman(nomor,nama,jenis,tanggal,lokasi,kondisi,tinggi,catatan) VALUES(?,?,?,?,?,?,?,?)",
                (f["nomor"].strip(),f["nama"].strip(),f.get("jenis",""),f.get("tanggal",""),f.get("lokasi",""),f.get("kondisi","Sehat"),f.get("tinggi",""),f.get("catatan","")))
            return redirect("/")
        except sqlite3.IntegrityError:
            return page("Nomor Duplikat","<div class='card'>Nomor sudah digunakan. <a href='/tambah'>Kembali</a></div>")
    with db() as c:
        n=c.execute("SELECT COALESCE(MAX(id),0)+1 FROM tanaman").fetchone()[0]
    return page("Tambah Tanaman",render_template_string(FORM,judul="Tambah Pohon",t=None,nomor=f"POHON-{n:03d}"))

@app.route("/tanaman/<int:id>")
def detail(id):
    with db() as c:
        t=c.execute("SELECT * FROM tanaman WHERE id=?",(id,)).fetchone()
    if not t: return "Tanaman tidak ditemukan",404
    body=render_template_string("""
    <div class="card"><h2>{{t['nama']}} — {{t['nomor']}}</h2>
    <img class="qr" src="/qr/{{t['id']}}">
    {% for label,key in [('Nomor','nomor'),('Nama','nama'),('Jenis','jenis'),('Tanggal Tanam','tanggal'),('Lokasi','lokasi'),('Kondisi','kondisi'),('Keterangan','catatan')] %}
    <p><b>{{label}}:</b> {{t[key] or '-'}}</p>{% endfor %}
    <div class="no-print">
    {% if session.get('admin') %}
      <a class="btn" href="/edit/{{t['id']}}">Edit</a>
    {% endif %}
    <button onclick="window.print()">Cetak QR</button>
    <a class="btn secondary" href="/">Kembali</a>
    {% if session.get('admin') %}
      <form method="post" action="/hapus/{{t['id']}}"
       onsubmit="return confirm('Hapus data ini?')">
       <button style="background:#b42318">Hapus</button>
      </form>
    {% endif %}
    </div></div>
    """,t=t)
    return page("Detail Tanaman",body)

@app.route("/qr/<int:id>")
def qr(id):
    with db() as c:
        t=c.execute("SELECT id FROM tanaman WHERE id=?",(id,)).fetchone()
    if not t: return "Tidak ditemukan",404
    alamat=request.host_url.rstrip("/")+url_for("detail",id=id)
    image=qrcode.make(alamat)
    out=io.BytesIO()
    image.save(out,format="PNG")
    out.seek(0)
    return send_file(out,mimetype="image/png")

@app.route("/edit/<int:id>",methods=["GET","POST"])
@admin_required
def edit(id):
    with db() as c:
        t=c.execute("SELECT * FROM tanaman WHERE id=?",(id,)).fetchone()
    if not t: return "Tanaman tidak ditemukan",404
    if request.method=="POST":
        f=request.form
        try:
            with db() as c:
                c.execute("UPDATE tanaman SET nomor=?,nama=?,jenis=?,tanggal=?,lokasi=?,kondisi=?,catatan=? WHERE id=?",
                (f["nomor"].strip(),f["nama"].strip(),f.get("jenis",""),f.get("tanggal",""),f.get("lokasi",""),f.get("kondisi","Sehat"),f.get("catatan",""),id))
            return redirect(url_for("detail",id=id))
        except sqlite3.IntegrityError:
            return page("Nomor Duplikat","<div class='card'>Nomor sudah dipakai pohon lain.</div>")
    return page("Edit Tanaman",render_template_string(FORM,judul="Edit Pohon",t=t,nomor=t["nomor"]))

@app.route("/hapus/<int:id>",methods=["POST"])
@admin_required
def hapus(id):
    with db() as c: c.execute("DELETE FROM tanaman WHERE id=?",(id,))
    return redirect("/")


@app.route("/export/excel")
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    with db() as c:
        rows = c.execute("""
            SELECT nomor, nama, jenis, tanggal, lokasi, kondisi, catatan
            FROM tanaman ORDER BY id DESC
        """).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Tanaman"
    ws.append([
        "Nomor Tanaman", "Nama Tanaman", "Jenis / Varietas",
        "Tanggal Tanam", "Lokasi / Blok", "Kondisi", "Keterangan"
    ])

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B45")
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([row[k] or "" for k in
                   ("nomor", "nama", "jenis", "tanggal",
                    "lokasi", "kondisi", "catatan")])

    widths = [20, 26, 22, 18, 22, 20, 40]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        as_attachment=True,
        download_name="laporan_data_tanaman.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route("/export/pdf")
def export_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from xml.sax.saxutils import escape

    with db() as c:
        rows = c.execute("""
            SELECT nomor, nama, jenis, tanggal, lokasi, kondisi
            FROM tanaman ORDER BY id DESC
        """).fetchall()

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=landscape(A4),
        rightMargin=12*mm, leftMargin=12*mm,
        topMargin=12*mm, bottomMargin=12*mm
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("LAPORAN DATA TANAMAN", styles["Title"]),
        Paragraph("Daftar identitas dan kondisi tanaman", styles["Normal"])
    ]

    data = [[
        "Nomor", "Nama Tanaman", "Jenis", "Tanggal",
        "Lokasi / Blok", "Kondisi"
    ]]

    for row in rows:
        data.append([
            Paragraph(escape(str(row[k] or "-")), styles["BodyText"])
            for k in ("nomor", "nama", "jenis", "tanggal", "lokasi", "kondisi")
        ])

    if len(data) == 1:
        data.append(["Belum ada data tanaman", "", "", "", "", ""])

    table = Table(
        data, repeatRows=1,
        colWidths=[32*mm, 48*mm, 38*mm, 27*mm, 48*mm, 35*mm]
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#176B45")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DCE6DF")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#F0F7F2")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(table)
    story.append(Paragraph(
        "Jumlah tanaman: " + str(len(rows)),
        styles["Normal"]
    ))
    doc.build(story)
    out.seek(0)

    return send_file(
        out,
        as_attachment=True,
        download_name="laporan_data_tanaman.pdf",
        mimetype="application/pdf"
    )


@app.route("/dashboard-lanjutan")
@admin_required
def dashboard_lanjutan():
    with db() as c:
        total = c.execute(
            "SELECT COUNT(*) FROM tanaman"
        ).fetchone()[0]

        kondisi_rows = c.execute("""
            SELECT COALESCE(NULLIF(TRIM(kondisi), ''), 'Belum diisi') AS kondisi,
                   COUNT(*) AS jumlah
            FROM tanaman
            GROUP BY COALESCE(NULLIF(TRIM(kondisi), ''), 'Belum diisi')
            ORDER BY jumlah DESC
        """).fetchall()

        lokasi_rows = c.execute("""
            SELECT COALESCE(NULLIF(TRIM(lokasi), ''), 'Belum diisi') AS lokasi,
                   COUNT(*) AS jumlah
            FROM tanaman
            GROUP BY COALESCE(NULLIF(TRIM(lokasi), ''), 'Belum diisi')
            ORDER BY jumlah DESC, lokasi
            LIMIT 10
        """).fetchall()

        terbaru = c.execute("""
            SELECT nomor, nama, lokasi, kondisi
            FROM tanaman
            ORDER BY id DESC
            LIMIT 8
        """).fetchall()

    warna = {
        "Sehat": "#16834a",
        "Perlu Perawatan": "#d97706",
        "Rusak": "#dc2626",
        "Mati": "#6b7280",
        "Belum diisi": "#94a3b8"
    }

    total_kondisi = sum(r["jumlah"] for r in kondisi_rows)
    kondisi_html = ""

    for r in kondisi_rows:
        kondisi = r["kondisi"]
        jumlah = r["jumlah"]
        persen = round(jumlah * 100 / total, 1) if total else 0
        warna_batang = warna.get(kondisi, "#176b43")

        kondisi_html += (
            '<div style="margin:18px 0">'
            '<div style="display:flex;justify-content:space-between;'
            'gap:12px;flex-wrap:wrap;margin-bottom:7px">'
            '<strong>' + kondisi + '</strong>'
            '<span>' + str(jumlah) + ' tanaman (' + str(persen) + '%)</span>'
            '</div>'
            '<div style="height:12px;background:#e8eee9;'
            'border-radius:20px;overflow:hidden">'
            '<div style="height:100%;width:' + str(persen) + '%;'
            'background:' + warna_batang + ';border-radius:20px"></div>'
            '</div></div>'
        )

    lokasi_html = ""
    for r in lokasi_rows:
        lokasi_html += (
            "<tr><td>" + str(r["lokasi"]) + "</td>"
            "<td><strong>" + str(r["jumlah"]) + "</strong></td></tr>"
        )

    terbaru_html = ""
    for r in terbaru:
        terbaru_html += (
            "<tr><td>" + str(r["nomor"] or "-") + "</td>"
            "<td>" + str(r["nama"] or "-") + "</td>"
            "<td>" + str(r["lokasi"] or "-") + "</td>"
            "<td>" + str(r["kondisi"] or "-") + "</td></tr>"
        )

    if not kondisi_html:
        kondisi_html = "<p>Belum ada data tanaman untuk dianalisis.</p>"
    if not lokasi_html:
        lokasi_html = '<tr><td colspan="2">Belum ada data lokasi.</td></tr>'
    if not terbaru_html:
        terbaru_html = '<tr><td colspan="4">Belum ada data tanaman.</td></tr>'

    body = """
    <div class="stats">
      <div class="stat">Total Tanaman<h2>""" + str(total) + """</h2></div>
      <div class="stat">Jenis Kondisi<h2>""" + str(len(kondisi_rows)) + """</h2></div>
      <div class="stat">Lokasi Terdata<h2>""" + str(len(lokasi_rows)) + """</h2></div>
    </div>

    <div class="card">
      <h2>Analisis Kondisi Tanaman</h2>
      <p>Distribusi kondisi berdasarkan data yang tersimpan.</p>
      """ + kondisi_html + """
    </div>

    <div class="card">
      <h2>Jumlah Tanaman per Lokasi</h2>
      <div class="scroll">
        <table>
          <thead><tr><th>Lokasi / Blok</th><th>Jumlah Tanaman</th></tr></thead>
          <tbody>""" + lokasi_html + """</tbody>
        </table>
      </div>
    </div>

    <div class="card">
      <h2>Data Tanaman Terbaru</h2>
      <div class="scroll">
        <table>
          <thead>
            <tr><th>Nomor</th><th>Nama</th><th>Lokasi</th><th>Kondisi</th></tr>
          </thead>
          <tbody>""" + terbaru_html + """</tbody>
        </table>
      </div>
      <p><a class="btn" href="/">Kembali ke Dashboard</a></p>
    </div>
    """

    return page("Analitik Tanaman", body)


if __name__=="__main__":
    app.run(host="0.0.0.0",port=5000,debug=False)


